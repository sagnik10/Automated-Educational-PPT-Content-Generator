import os
import re
import glob
import hashlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

ASSETS_TOPIC_DIR = "assets_topic_images"
ASSETS_REF_DIR = "images_export"
OUT_FILE = "MASTER_FINAL_SAMPLESTYLE_LOGO_SCANNER_STORY.pptx"

PAT_REMOVE = re.compile(
    r"\bSlow\s*Learners?\b|\bSlow\s*Steps?\b|\bUse\s*slow\s*steps?\b|\bSolve\s*slowly\b|\bGo\s*slowly\b|\bWe\s*practice\s*slowly\b",
    re.IGNORECASE,
)

CANDIDATE_PPTS = [
    "RATIOS_G7_UnitRates_Fractions_PAGES1to5_SMOOTH_HEADERS_FIXED.pptx",
    "RATIOS_G7_Lesson3_2_ConnectPercent_Styled_Kids_NoOverlap.pptx",
    "Lesson3_2_and_3_Styled_Kids_NoOverlap_Pages152to157_FIXED.pptx",
    "Lesson3_3_Pages157to160_Long_Kids_NoOverlap_FIXED.pptx",
]

def clean_text(s):
    if not s:
        return s
    s = PAT_REMOVE.sub("", s)
    s = re.sub(r"\bThen\s*\.\s*", "", s, flags=re.IGNORECASE)
    s = re.sub(r"\bThen\s*$", "", s, flags=re.IGNORECASE)
    s = re.sub(r"\(\s*\)", "", s)
    s = re.sub(r"\s{2,}", " ", s)
    s = re.sub(r"\s+\n", "\n", s)
    s = re.sub(r"\n\s+", "\n", s)
    return s.strip()

def clean_textframe(tf):
    for p in tf.paragraphs:
        for r in p.runs:
            r.text = clean_text(r.text)

def clean_slide_text(slide):
    for sh in slide.shapes:
        if getattr(sh, "has_text_frame", False):
            clean_textframe(sh.text_frame)
        if getattr(sh, "has_table", False):
            for row in sh.table.rows:
                for cell in row.cells:
                    if cell.text_frame:
                        clean_textframe(cell.text_frame)

def slide_all_text(slide):
    parts = []
    for sh in slide.shapes:
        if getattr(sh, "has_text_frame", False):
            parts.append(sh.text_frame.text or "")
        if getattr(sh, "has_table", False):
            for row in sh.table.rows:
                for cell in row.cells:
                    if cell.text_frame:
                        parts.append(cell.text_frame.text or "")
    return "\n".join([p for p in parts if p])

def page_num(text):
    m = re.search(r"\bPage\s+(\d+)\b", text, flags=re.IGNORECASE)
    if m:
        try:
            return int(m.group(1))
        except Exception:
            return None
    return None

def is_cover(text):
    t = re.sub(r"\s+", " ", (text or "").strip().lower())
    return ("ratios" in t and "grade 7" in t and len(t) <= 90)

def sig(text):
    t = clean_text(text)
    t = re.sub(r"\bPage\s+\d+\b", "", t, flags=re.IGNORECASE)
    t = re.sub(r"\s{2,}", " ", t).strip().lower()
    return hashlib.md5(t.encode("utf-8", errors="ignore")).hexdigest()

def rr(slide, x, y, w, h, fill, line=None, line_pt=0, rounded=True):
    typ = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    sh = slide.shapes.add_shape(typ, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line and line_pt > 0:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_pt)
    else:
        sh.line.fill.background()
    return sh

def list_images(dirpath):
    if not os.path.isdir(dirpath):
        return []
    exts = (".png", ".jpg", ".jpeg", ".webp")
    files = []
    for p in glob.glob(os.path.join(dirpath, "**", "*.*"), recursive=True):
        if os.path.splitext(p)[1].lower() in exts:
            files.append(p)
    files.sort()
    return files

def pick_logo_and_scanner(ref_imgs):
    logo = None
    scan = None
    for p in ref_imgs:
        b = os.path.basename(p).lower()
        if logo is None and ("logo" in b or "ib" in b or "continuum" in b):
            logo = p
        if scan is None and ("qr" in b or "scan" in b or "scanner" in b or "barcode" in b):
            scan = p
    if logo is None and ref_imgs:
        logo = ref_imgs[0]
    if scan is None and ref_imgs:
        scan = ref_imgs[-1]
    return logo, scan

def add_scanner_bar(slide, scan_img, sw, sh):
    if not scan_img or not os.path.isfile(scan_img):
        return
    bar_h = Inches(0.55)
    y = sh - bar_h
    rr(slide, 0, y, sw, bar_h, RGBColor(255, 255, 255), RGBColor(201, 221, 242), 1.2, False)
    try:
        slide.shapes.add_picture(scan_img, sw - Inches(1.65), y + Inches(0.05), Inches(1.55), Inches(0.45))
    except Exception:
        pass
    tb = slide.shapes.add_textbox(Inches(0.60), y + Inches(0.13), sw - Inches(2.40), Inches(0.30))
    tf = tb.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "Scan here ✅"
    p.alignment = PP_ALIGN.LEFT
    r = p.runs[0]
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(12, 34, 56)

def logo_full_slide(prs, logo_img, scanner_img):
    sw, sh = prs.slide_width, prs.slide_height
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rr(s, 0, 0, sw, sh, RGBColor(255, 255, 255), None, 0, False)
    if logo_img and os.path.isfile(logo_img):
        try:
            s.shapes.add_picture(logo_img, 0, 0, sw, sh)
        except Exception:
            try:
                s.shapes.add_picture(logo_img, Inches(1.0), Inches(1.0), sw - Inches(2.0), sh - Inches(2.0))
            except Exception:
                pass
    add_scanner_bar(s, scanner_img, sw, sh)
    return s

def header(slide, title, subtitle, sw, sh):
    rr(slide, 0, 0, sw, sh, RGBColor(245, 250, 255), None, 0, False)
    rr(slide, Inches(0.60), Inches(0.28), sw - Inches(1.20), Inches(1.20),
       RGBColor(232, 245, 233), RGBColor(46, 125, 50), 2, True)

    tb = slide.shapes.add_textbox(Inches(0.95), Inches(0.50), sw - Inches(1.90), Inches(0.60))
    tf = tb.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.10)
    tf.margin_right = Inches(0.10)
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    r = p.runs[0]
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = RGBColor(11, 43, 22)

    tb2 = slide.shapes.add_textbox(Inches(0.95), Inches(1.12), sw - Inches(1.90), Inches(0.40))
    tf2 = tb2.text_frame
    tf2.clear()
    tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf2.margin_left = Inches(0.10)
    tf2.margin_right = Inches(0.10)
    p2 = tf2.paragraphs[0]
    p2.text = subtitle
    p2.alignment = PP_ALIGN.LEFT
    r2 = p2.runs[0]
    r2.font.size = Pt(15)
    r2.font.color.rgb = RGBColor(12, 34, 56)

def story_slide(prs, title, subtitle, big_img, story_lines, problems_lines, steps_lines, scanner_img):
    sw, sh = prs.slide_width, prs.slide_height
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, title, subtitle, sw, sh)

    rr(s, Inches(0.60), Inches(1.70), sw - Inches(1.20), sh - Inches(2.75),
       RGBColor(255, 255, 255), RGBColor(201, 221, 242), 2, True)

    img_x = Inches(0.95)
    img_y = Inches(2.05)
    img_w = Inches(5.10)
    img_h = Inches(4.35)

    rr(s, img_x - Inches(0.10), img_y - Inches(0.10), img_w + Inches(0.20), img_h + Inches(0.20),
       RGBColor(255, 255, 255), RGBColor(201, 221, 242), 1.6, True)
    if big_img and os.path.isfile(big_img):
        try:
            s.shapes.add_picture(big_img, img_x, img_y, img_w, img_h)
        except Exception:
            pass

    right_x = Inches(6.20)
    right_y = Inches(2.05)
    right_w = sw - right_x - Inches(0.95)
    right_h = Inches(4.35)

    rr(s, right_x, right_y, right_w, right_h, RGBColor(245, 250, 255), RGBColor(201, 221, 242), 1.6, True)

    def block(y, h, title_txt, lines, accent):
        rr(s, right_x + Inches(0.25), y, right_w - Inches(0.50), h,
           RGBColor(255, 255, 255), RGBColor(201, 221, 242), 1.2, True)
        rr(s, right_x + Inches(0.25), y, right_w - Inches(0.50), Inches(0.45),
           accent, RGBColor(201, 221, 242), 1.0, True)
        tb = s.shapes.add_textbox(right_x + Inches(0.40), y + Inches(0.06), right_w - Inches(0.80), Inches(0.35))
        tf = tb.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = title_txt
        p.alignment = PP_ALIGN.LEFT
        r = p.runs[0]
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = RGBColor(12, 34, 56)

        tb2 = s.shapes.add_textbox(right_x + Inches(0.40), y + Inches(0.55), right_w - Inches(0.80), h - Inches(0.65))
        tf2 = tb2.text_frame
        tf2.clear()
        tf2.word_wrap = True
        tf2.vertical_anchor = MSO_ANCHOR.TOP
        tf2.margin_left = Inches(0.06)
        tf2.margin_right = Inches(0.06)
        for i, ln in enumerate(lines):
            p2 = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
            p2.text = ln
            p2.alignment = PP_ALIGN.LEFT
            run = p2.runs[0]
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(12, 34, 56)

    block(right_y + Inches(0.20), Inches(1.35), "📖 Story", story_lines, RGBColor(232, 245, 233))
    block(right_y + Inches(1.70), Inches(1.45), "🧩 Problems", problems_lines, RGBColor(255, 243, 196))
    block(right_y + Inches(3.30), Inches(0.95), "✅ Steps", steps_lines, RGBColor(238, 246, 255))

    add_scanner_bar(s, scanner_img, sw, sh)
    return s

def resolve_sources(folder):
    existing = {f.lower(): f for f in os.listdir(folder) if f.lower().endswith(".pptx")}
    out = []
    for want in CANDIDATE_PPTS:
        if want.lower() in existing:
            out.append(os.path.join(folder, existing[want.lower()]))
            continue
        base = os.path.splitext(want)[0].lower()
        best = None
        for k, real in existing.items():
            if base in k:
                best = real
                break
        if best:
            out.append(os.path.join(folder, best))
    if len(out) < 2:
        raise FileNotFoundError("Not enough source PPTX files found in the folder. Check filenames in PPT Work.")
    return out

def build():
    folder = os.path.dirname(os.path.abspath(__file__))
    src_paths = resolve_sources(folder)
    topic_root = os.path.join(folder, ASSETS_TOPIC_DIR)
    ref_root = os.path.join(folder, ASSETS_REF_DIR)
    out_path = os.path.join(folder, OUT_FILE)

    ref_imgs = list_images(ref_root)
    logo_img, scanner_img = pick_logo_and_scanner(ref_imgs)

    base = Presentation(src_paths[0])
    master = Presentation()
    master.slide_width = base.slide_width
    master.slide_height = base.slide_height
    sw, sh = master.slide_width, master.slide_height

    logo_full_slide(master, logo_img, scanner_img)

    raw = []
    for p in src_paths:
        prs = Presentation(p)
        for sl in prs.slides:
            txt = slide_all_text(sl)
            pg = page_num(txt)
            cov = is_cover(txt)
            raw.append((cov, pg if pg is not None else 10**9, txt, sl))

    by_sig = {}
    keep_cover = True
    for cov, pg, txt, sl in raw:
        if cov:
            if keep_cover:
                keep_cover = False
            else:
                continue
        k = sig(txt)
        if k not in by_sig:
            by_sig[k] = (cov, pg, txt, sl)
        else:
            _, pg0, _, _ = by_sig[k]
            if pg0 == 10**9 and pg != 10**9:
                by_sig[k] = (cov, pg, txt, sl)

    uniq = list(by_sig.values())
    uniq.sort(key=lambda x: (0 if x[0] else 1, x[1]))

    def copy_slide_into(src_slide):
        dest = master.slides.add_slide(master.slide_layouts[6])
        for shp in list(dest.shapes):
            try:
                dest.shapes._spTree.remove(shp._element)
            except Exception:
                pass
        for shp in src_slide.shapes:
            dest.shapes._spTree.insert_element_before(shp._element, "p:extLst")
        clean_slide_text(dest)
        add_scanner_bar(dest, scanner_img, sw, sh)
        return dest

    unit_liters = os.path.join(topic_root, "unit_rates", "liters_day.png")
    unit_speed = os.path.join(topic_root, "unit_rates", "speed_kmh.png")
    frac_pizza = os.path.join(topic_root, "fractions", "fraction_pizza.png")
    pp_grid = os.path.join(topic_root, "percent_proportion", "percent_100_grid.png")
    rl_receipt = os.path.join(topic_root, "real_life", "receipt_tax.png")
    rl_commission = os.path.join(topic_root, "real_life", "commission_house.png")
    rl_shoes = os.path.join(topic_root, "real_life", "shoes_compare.png")

    inserted_unit = False
    inserted_pp = False
    inserted_pe = False

    for cov, pg, txt, sl in uniq:
        copy_slide_into(sl)

        if (not inserted_unit) and pg <= 5:
            story_slide(master, "Story: Water for the Class 💧", "Unit rate = liters per 1 day",
                        unit_liters,
                        ["The class drinks 12 liters in 4 days.", "How many liters per 1 day?"],
                        ["1) Find L/day.", "2) For 7 days, how many liters?"],
                        ["Unit rate = total ÷ days", "Then multiply by 7"],
                        scanner_img)
            story_slide(master, "Story: Bus Trip 🚍", "Unit rate = km per 1 hour",
                        unit_speed,
                        ["A bus goes 180 km in 3 hours.", "How many km per 1 hour?"],
                        ["1) Find km/h.", "2) In 5 hours, how far?"],
                        ["Unit rate = distance ÷ time", "Then multiply"],
                        scanner_img)
            story_slide(master, "Story: Pizza Fractions 🍕", "Fractions show parts of a whole",
                        frac_pizza,
                        ["A pizza has 8 equal slices.", "Sam eats 3 slices."],
                        ["1) Write the fraction eaten.", "2) How many slices left?"],
                        ["Fraction = part/whole", "Left = whole − part"],
                        scanner_img)
            inserted_unit = True

        if (not inserted_pp) and pg >= 149:
            story_slide(master, "Story: Sticker Chart 🟩", "Percent = shaded out of 100",
                        pp_grid,
                        ["A 100-grid has 100 squares.", "You shade 37 squares."],
                        ["1) What percent is shaded?", "2) What decimal is that?"],
                        ["Percent = shaded/100", "Decimal = percent ÷ 100"],
                        scanner_img)
            inserted_pp = True

        if (not inserted_pe) and pg >= 155:
            story_slide(master, "Story: Tax on a Receipt 🧾", "part = percent × whole",
                        rl_receipt,
                        ["Bill is $54 and tax is 8.44%.", "Find the tax amount."],
                        ["1) Write the equation.", "2) Compute/estimate tax."],
                        ["Convert % to decimal", "Multiply decimal × whole"],
                        scanner_img)
            story_slide(master, "Story: House Commission 🏠", "Find the whole (selling price)",
                        rl_commission,
                        ["Commission is 5.5% and equals $9,020.", "Find the selling price."],
                        ["1) 9020 = 0.055 × whole", "2) whole = 9020 ÷ 0.055"],
                        ["Whole = part ÷ percent", "Check by multiplying"],
                        scanner_img)
            story_slide(master, "Story: Shoe Reviews 👟", "Compare percents to decide",
                        rl_shoes,
                        ["Shoe A: 42/50 good reviews.", "Shoe B: 78/100 good reviews."],
                        ["1) Turn both into percents.", "2) Pick the better shoe and explain."],
                        ["Percent = part/whole × 100", "Compare the percents"],
                        scanner_img)
            inserted_pe = True

    master.save(out_path)
    print(out_path)

if __name__ == "__main__":
    build()